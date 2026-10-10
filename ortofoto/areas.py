"""Cierre de areas (veredas, cunetas, pistas, martillos) y metrados.

Franja: un elemento con dos bordes levantados (p.ej. vereda: borde de
fachada/lote y borde de sardinel, ambos con codigo VER). Cada muestra de un
borde "vota" por el borde del mismo codigo mas cercano si esta a una distancia
entre ancho_min y ancho_max. Dos bordes que se votan mutuamente forman una
franja; el area se cierra solo en el tramo donde realmente son paralelos, asi
una vereda no se pega con el tramo de la otra cuadra.

Contorno: un borde que ya se cerro sobre si mismo (martillo, acceso).

Las areas no se superponen: cada nueva area se recorta contra las ya aceptadas.
"""
from dataclasses import dataclass, field

import numpy as np
from scipy.spatial import cKDTree
from shapely import STRtree
from shapely.geometry import LineString, MultiPolygon, Point, Polygon
from shapely.ops import substring, unary_union

PASO = 0.5  # m entre muestras de un borde
TRAMO_MIN = 1.0  # m minimos de borde paralelo para cerrar un area
CORTE_TRAMO = 2.0  # m sin pareja que separan un tramo en dos areas
AREA_MIN = 0.15  # m2


@dataclass
class Area:
    codigo: str
    conf: object
    poligono: Polygon
    origen: str  # franja | contorno
    ancho: float = 0.0
    largo: float = 0.0
    revisar: bool = False
    etiqueta: str = ""
    nota: str = ""
    lamina: str = ""
    reglas: list = field(default_factory=list)  # criterios que formaron o recortaron el area

    @property
    def area(self):
        return float(self.poligono.area)


@dataclass
class Linea:
    codigo: str
    conf: object
    linea: LineString
    etiqueta: str = ""
    sin_pareja: bool = False  # borde de franja que no se pudo cerrar
    borde: bool = False  # borde de un area armada por secciones (se dibuja en la capa del codigo)
    dibujar: bool = False  # la linea no viene de los puntos (corte lineal): se dibuja en su capa
    lamina: str = ""


@dataclass
class Metrado:
    areas: list = field(default_factory=list)
    lineas: list = field(default_factory=list)
    puntos_revisar: list = field(default_factory=list)  # (x, y, motivo)
    sin_puntos: list = field(default_factory=list)  # concreto visto en la foto sin respaldo de puntos


def _bordes(res):
    abiertos, cerrados = [], []
    for cad in res.cadenas:
        xy = res.xy(cad)
        if cad[0] == cad[-1] and len(cad) >= 4:
            cerrados.append(xy)
        elif len(cad) >= 2:
            abiertos.append(LineString(xy))
    return abiertos, cerrados


def _tramos(posiciones):
    """Agrupa posiciones (m a lo largo del borde) en tramos continuos."""
    posiciones = np.sort(posiciones)
    cortes = np.where(np.diff(posiciones) > CORTE_TRAMO)[0]
    inicio = 0
    for c in list(cortes) + [len(posiciones) - 1]:
        yield posiciones[inicio], posiciones[c]
        inicio = c + 1


def _franjas(res, bordes):
    conf = res.conf
    if len(bordes) < 2:
        return [], set()
    arbol = STRtree(bordes)
    votos = {}  # (a, b) -> lista de (pos_en_a, distancia)
    for ia, a in enumerate(bordes):
        n = max(2, int(a.length / PASO) + 1)
        for pos in np.linspace(0, a.length, n):
            p = a.interpolate(pos)
            cercanos = arbol.query(p.buffer(conf.ancho_max))
            mejor = None
            for ib in cercanos:
                if ib == ia:
                    continue
                b = bordes[ib]
                d = b.distance(p)
                if mejor is None or d < mejor[1]:
                    mejor = (ib, d)
            if mejor and conf.ancho_min <= mejor[1] <= conf.ancho_max:
                b = bordes[mejor[0]]
                pb = b.project(p)
                # La pareja debe estar "enfrente", no pasado el extremo del otro borde
                if 0.05 < pb < b.length - 0.05 or b.distance(p) < conf.ancho_max * 0.8:
                    votos.setdefault((ia, mejor[0]), []).append((pos, mejor[1]))
    areas, usados = [], set()
    for (ia, ib), vs in votos.items():
        if ia > ib and (ib, ia) in votos:
            continue  # cada pareja una sola vez
        if (ib, ia) not in votos:
            continue  # el otro borde no la reconoce como pareja
        a, b = bordes[ia], bordes[ib]
        pos = np.array([v[0] for v in vs])
        for p0, p1 in _tramos(pos):
            if p1 - p0 < TRAMO_MIN:
                continue
            ta = substring(a, p0, p1)
            proy = [b.project(a.interpolate(x)) for x in np.linspace(p0, p1, 8)]
            q0, q1 = min(proy), max(proy)
            if q1 - q0 < TRAMO_MIN * 0.5:
                continue
            tb = substring(b, q0, q1)
            ca, cb = list(ta.coords), list(tb.coords)
            # Cerrar sin cruzar: el final de A se une con el extremo mas cercano de B
            if np.hypot(*np.subtract(ca[-1], cb[0])) > np.hypot(*np.subtract(ca[-1], cb[-1])):
                cb = cb[::-1]
            pol = Polygon(ca + cb)
            if not pol.is_valid:
                pol = pol.buffer(0)
                if isinstance(pol, MultiPolygon):
                    pol = max(pol.geoms, key=lambda g: g.area)
            if pol.is_empty or pol.area < AREA_MIN:
                continue
            ancho = float(np.median([v[1] for v in vs if p0 <= v[0] <= p1]))
            largo = (ta.length + tb.length) / 2
            # Forma coherente con un ancho y un largo; si no, a revisar
            revisar = not (0.6 < pol.area / max(ancho * largo, 1e-6) < 1.5)
            areas.append(Area(res.codigo, conf, pol, "franja", ancho, largo, revisar))
            usados.update((ia, ib))
    return areas, usados


def _poligono_valido(coords):
    pol = Polygon(coords)
    if not pol.is_valid:
        pol = pol.buffer(0)
        if isinstance(pol, MultiPolygon):
            pol = max(pol.geoms, key=lambda g: g.area)
    return pol


def cerrar_areas(resultados, base=None, usar_foto=False):
    """Devuelve Metrado con areas (sin superposiciones) y lineas.

    base: PlanoBase con los limites de propiedad (fachadas) del plano del proyecto.
    """
    con_base = base is not None and not base.vacio
    met = Metrado()
    candidatas = []
    lineas_por_codigo = {}
    for res in resultados:
        lineas_por_codigo[res.codigo] = [LineString(res.xy(c)) for c in res.cadenas if len(c) >= 2]
    for res in resultados:
        conf = res.conf
        if conf.tipo == "franja" and conf.referencia:
            from .veredas import veredas_contra_fachada

            refs = [ln for cod in conf.referencia for ln in lineas_por_codigo.get(cod, [])]
            csh = [(p.e, p.n) for r in resultados if r.codigo in conf.referencia for p in r.puntos]
            ars, bordes, usados = [], [], set()
            if con_base:
                # El borde interior es el limite de propiedad dibujado en el plano
                ars, bordes, usados = veredas_contra_fachada(
                    res, base.limites, base.union_manzanas() if base.manzanas else None, csh,
                    revisar=met.puntos_revisar)
                if usar_foto and res.orto is not None:
                    ars = _ajustar_con_foto(res, base, ars, met)
            # Puntos sin limite del plano a su alcance (la fachada del plano esta bajo el techo, o
            # corrida hacia la calle, o no hay plano base): contra la fachada levantada (CSH, LP)
            resto = [i for i in range(len(res.puntos)) if i not in usados]
            if refs and resto:
                a_, b_, u_ = veredas_contra_fachada(res, refs, None, csh, indices=resto,
                                                        revisar=met.puntos_revisar)
                for a in a_:
                    a.origen = "fachada levantada"  # ya esta pegada a la fachada real: no se recorta
                ars, bordes, usados = ars + a_, bordes + b_, usados | u_
            resto = [i for i in range(len(res.puntos)) if i not in usados]
            ars2, bordes2, usados2 = _franjas_por_estaciones(res, resto, res.barreras)
            candidatas += ars + ars2
            met.lineas += [Linea(res.codigo, conf, b, borde=True) for b in bordes + bordes2]
            _unir_sobrantes(res, usados | usados2)
        elif conf.tipo == "franja":
            # Franja sin referencia (canal, cuneta, pista):
            # 1) figuras cerradas levantadas por sus esquinas (cajas de alcantarilla)
            ars, usados = _figuras_cerradas(res) if conf.ancho_defecto else ([], set())
            # 2) secciones a lo ancho (dos bordes cada pocos metros), como las veredas
            resto = [i for i in range(len(res.puntos)) if i not in usados]
            ars2, bordes, usados2 = _franjas_por_estaciones(res, resto, res.barreras)
            ars += ars2
            usados |= usados2
            # 3) lo demas se une a lo largo (eje) y, si hay ancho supuesto, se cierra con ese ancho
            _unir_sobrantes(res, usados)
            if conf.ancho_defecto:
                ars3 = _franjas_por_eje(res)
                ars += ars3
                if ars3:
                    met.lineas += [Linea(res.codigo, conf, LineString(res.xy(c)), borde=True) for c in res.cadenas]
                    res.cadenas = []
            if usar_foto and res.orto is not None:
                ars = _cortar_con_foto(res.orto, ars)
            candidatas += ars
            met.lineas += [Linea(res.codigo, conf, b, borde=True) for b in bordes]
        if conf.tipo == "punto" or not res.cadenas:
            continue
        abiertos, cerrados = _bordes(res)
        if conf.tipo == "contorno" and conf.capa_area:
            # Contorno abierto (p.ej. martillo sin el lado del sardinel): cerrar si los extremos estan cerca
            quedan = []
            for ln in abiertos:
                c = list(ln.coords)
                if len(c) >= 3 and np.hypot(*np.subtract(c[0], c[-1])) <= conf.separacion_max:
                    cerrados.append(np.array(c + [c[0]]))
                else:
                    quedan.append(ln)
            abiertos = quedan
        if conf.tipo in ("franja", "contorno") and conf.capa_area:
            for xy in cerrados:
                pol = Polygon(xy)
                if not pol.is_valid:
                    pol = pol.buffer(0)
                    if isinstance(pol, MultiPolygon):
                        pol = max(pol.geoms, key=lambda g: g.area)
                if pol.area >= AREA_MIN:
                    candidatas.append(Area(res.codigo, conf, pol, "contorno", largo=pol.length / 2))
            usados = set()
            if conf.tipo == "franja":
                areas, usados = _franjas(res, abiertos)
                candidatas += areas
            for k, ln in enumerate(abiertos):
                met.lineas.append(Linea(res.codigo, conf, ln, sin_pareja=(conf.tipo == "franja" and k not in usados)))
        else:
            met.lineas += [Linea(res.codigo, conf, ln) for ln in abiertos]
            met.lineas += [Linea(res.codigo, conf, LineString(xy)) for xy in cerrados]

    if con_base:
        # Ningun area entra a los lotes
        recortadas = []
        for ar in candidatas:
            if ar.origen in ("fachada", "fachada levantada"):  # ya recortadas contra los lotes
                recortadas.append(ar)
                continue
            pol = base.recortar(ar.poligono)
            if isinstance(pol, Polygon) and pol.area >= AREA_MIN:
                ar.poligono = pol
                recortadas.append(ar)
        candidatas = recortadas
    # Sin superposiciones: primero las mas confiables (contornos, luego franjas grandes)
    # Los canales y cunetas van primero: la vereda se corta contra ellos, no los pisa
    candidatas.sort(key=lambda a: (bool(a.conf.referencia), a.revisar, a.origen != "contorno", -a.area))
    ocupado = None
    for ar in candidatas:
        pol = ar.poligono
        if ocupado is not None and pol.intersects(ocupado):
            resto = pol.difference(ocupado)
            if resto.area < 0.5 * pol.area:
                continue  # es practicamente un duplicado
            if isinstance(resto, MultiPolygon):
                resto = max(resto.geoms, key=lambda g: g.area)
            if resto.area < AREA_MIN or not isinstance(resto, Polygon):
                continue
            ar.poligono = resto
            if ar.conf.referencia:
                ar.reglas.append("recortada contra canal/cuneta (tienen prioridad)")
        met.areas.append(ar)
        ocupado = ar.poligono if ocupado is None else unary_union([ocupado, ar.poligono])
    _numerar(met)
    return met


DESNIVEL_MAX = 0.30  # m de salto de cota entre dos puntos VER seguidos
PENDIENTE_MAX = 0.25  # y con mas pendiente que esto (una vereda baja 2-5 % hacia la pista)
VEGETACION_MIN = 0.6  # fraccion visible del tramo entre los dos puntos con vegetacion (jardin)


def _continua(a, b, dd, zz, xy, orto):
    """Motivo por el que el punto b (mas afuera) ya no es la misma vereda que a, o "" si lo es.

    Mandan los puntos: solo se corta en un caso claro, un jardin. Tiene que haber un desnivel
    brusco Y vegetacion a la vista en la mayor parte del tramo entre los dos puntos. Bajo la copa
    de un arbol la foto ve sombra (no decide) y no se corta; el polvo sobre el concreto tampoco.
    """
    if orto is None or xy is None or zz is None:
        return ""
    dz = abs(zz[b] - zz[a])
    paso = max(np.hypot(*np.subtract(xy[b], xy[a])), 0.05)
    if dz <= DESNIVEL_MAX or dz / paso <= PENDIENTE_MAX:
        return ""
    f = _fracciones_tramo(orto, xy[a], xy[b])
    if f is not None and f["vegetacion"] >= VEGETACION_MIN:
        return f"jardin: vegetacion en la foto y desnivel de {dz:.2f} m"
    return ""


def _fracciones_tramo(orto, p, q, ancho=0.3):
    from .concreto import fracciones_suelo

    ln = LineString([tuple(p), tuple(q)])
    if ln.length < 0.2:
        return None
    try:
        return fracciones_suelo(orto, ln.buffer(ancho / 2, cap_style=2))
    except Exception:  # la foto no alcanza: no se decide con la foto
        return None


FIGURA_ENLACE = 2.5  # m: puntos mas cercanos que esto forman la misma figura
FIGURA_MAX = 8.0  # m: una figura cerrada (caja) no mide mas que esto
FIGURA_ANCHO_MIN = 0.4  # m: lado menor minimo para que el grupo sea una figura y no una linea


def _figuras_cerradas(res):
    """Grupos chicos de 3 o mas puntos que forman una figura (no una linea): se cierran con su forma.

    El contorno recorre los puntos en orden alrededor de su centro (no la envolvente), asi respeta
    la figura levantada. Devuelve (areas, indices_usados).
    """
    from scipy.sparse import coo_matrix
    from scipy.sparse.csgraph import connected_components
    from shapely.geometry import MultiPoint

    conf = res.conf
    xy = np.array([(p.e, p.n) for p in res.puntos])
    if len(xy) < 3:
        return [], set()
    pares = np.array(list(cKDTree(xy).query_pairs(FIGURA_ENLACE)) or [(0, 0)])
    n, lab = connected_components(coo_matrix((np.ones(len(pares)), (pares[:, 0], pares[:, 1])),
                                             shape=(len(xy), len(xy))), directed=False)
    areas, usados = [], set()
    for k in range(n):
        idx = np.where(lab == k)[0]
        if not 3 <= len(idx) <= 12:
            continue
        m = xy[idx]
        rect = MultiPoint([tuple(q) for q in m]).minimum_rotated_rectangle
        if not isinstance(rect, Polygon):
            continue
        c = np.array(rect.exterior.coords)[:4]
        lados = sorted(np.hypot(*(c[1:] - c[:-1]).T).tolist() + [float(np.hypot(*(c[0] - c[3])))])
        if lados[0] < FIGURA_ANCHO_MIN or lados[-1] > FIGURA_MAX:
            continue
        centro = m.mean(axis=0)
        orden = np.argsort(np.arctan2(m[:, 1] - centro[1], m[:, 0] - centro[0]))
        pol = _poligono_valido([tuple(q) for q in m[orden]])
        if pol.is_empty or pol.area < AREA_MIN:
            continue
        areas.append(Area(res.codigo, conf, pol, "figura", lados[0], lados[-1], nota="figura cerrada por sus puntos"))
        usados.update(int(i) for i in idx)
    return areas, usados


def _franjas_por_eje(res):
    """Tramos levantados solo por su eje (pares inicio/fin, cadenas a lo largo): se cierran con el
    ancho supuesto del codigo. El largo es el del eje."""
    conf = res.conf
    w = conf.ancho_defecto
    areas = []
    for cad in res.cadenas:
        if len(cad) < 2:
            continue
        eje = LineString(res.xy(cad))
        if eje.length < 0.3:
            continue
        pol = eje.buffer(w / 2, cap_style=2, join_style=2)
        if isinstance(pol, MultiPolygon):
            pol = max(pol.geoms, key=lambda g: g.area)
        areas.append(Area(res.codigo, conf, pol, "eje", w, eje.length,
                          nota=f"levantado por su eje: ancho supuesto {w:.2f} m"))
    return areas


def _cortar_con_foto(orto, ars):
    """Corta cada area donde la foto muestra tierra o pasto atravesandola (cada pedazo queda cerrado)."""
    from .concreto import cortar_por_foto

    salida = []
    for a in ars:
        piezas, quitado = cortar_por_foto(orto, a.poligono)
        if quitado < 0.3 or not piezas:
            salida.append(a)
            continue
        for g in piezas:
            salida.append(Area(a.codigo, a.conf, g, a.origen, a.ancho, a.largo * g.area / max(a.area, 1e-6),
                               a.revisar, nota=f"cortada donde la foto no muestra concreto ({quitado:.1f} m2 quitados)"))
    return salida


def _ajustar_con_foto(res, base, ars_puntos, met):
    """Reemplaza la forma de cada vereda por la que se ve en la ortofoto, si coincide con los puntos."""
    from .concreto import veredas_desde_foto

    conf = res.conf
    xy = np.array([(p.e, p.n) for p in res.puntos])
    grupos = {}
    for i, (x, y) in enumerate(xy):
        k = base.limite_de(x, y, dist=conf.ancho_max + 1.0)
        if k >= 0:
            grupos.setdefault(k, []).append(i)
    vistas = veredas_desde_foto(res.orto, base, xy, conf, grupos)
    for v in vistas:
        for i in v.lejos:
            met.puntos_revisar.append((float(xy[i][0]), float(xy[i][1]),
                                       f"{res.puntos[i].num} {res.puntos[i].desc}: no calza con la foto"))
    final, usadas = [], set()
    for a in ars_puntos:
        cand = [k for k, v in enumerate(vistas) if v.poligono.intersects(a.poligono)]
        if not cand:
            # Bajo alero o en sombra: la foto no muestra esta vereda; queda la forma de los puntos
            a.nota = "no visible en la foto (alero/sombra): forma de los puntos"
            final.append(a)
            continue
        foto = unary_union([vistas[k].poligono for k in cand]).intersection(a.poligono.buffer(1.0))
        foto = base.recortar(foto) if not foto.is_empty else foto
        if isinstance(foto, MultiPolygon):
            foto = max(foto.geoms, key=lambda g: g.area)
        relacion = foto.area / max(a.area, 1e-6) if isinstance(foto, Polygon) else 0
        if isinstance(foto, Polygon) and 0.5 <= relacion <= 2.0:
            lejos = any(i in vistas[k].lejos for k in cand for i in vistas[k].puntos
                        if a.poligono.buffer(0.5).contains(Point(xy[i])))
            a.poligono, a.origen = foto, "foto"
            a.revisar = a.revisar or lejos
            a.nota = "forma de la foto" + ("; hay puntos que no calzan" if lejos else "")
            usadas.update(cand)
        else:
            a.revisar, a.nota = True, f"la foto no coincide con los puntos ({relacion:.0%} del area)"
        final.append(a)
    # Donde la foto muestra tierra o pasto atravesando la vereda, no hay concreto: se corta
    from .concreto import cortar_por_foto

    cortadas = []
    for a in final:
        piezas, quitado = cortar_por_foto(res.orto, a.poligono)
        if quitado < 0.5 or not piezas:
            cortadas.append(a)
            continue
        for g in piezas:
            cortadas.append(Area(a.codigo, a.conf, g, a.origen, a.ancho, a.largo, a.revisar,
                                 nota=(a.nota + "; " if a.nota else "") +
                                 f"cortada donde la foto no muestra concreto ({quitado:.1f} m2 quitados)"))
    final = cortadas
    # Concreto visto en la foto, pegado a la fachada, que no salio de los puntos
    for k, v in enumerate(vistas):
        if k not in usadas and not any(v.poligono.intersects(a.poligono) for a in final):
            final.append(Area(res.codigo, conf, v.poligono, "foto", revisar=True,
                              nota="vista en la foto, sin area de puntos"))
    return final


def _franjas_por_estaciones(res, indices, barreras=None):
    """Franjas levantadas por secciones sin linea de referencia cerca.

    1) Seccion = grupo de 2 a 5 puntos vecinos cuyo ancho (los dos mas
       alejados) esta entre ancho_min y ancho_max.
    2) Se encadenan secciones vecinas (hasta separacion_max) solo si se avanza
       de forma perpendicular a ambas secciones (una vereda avanza a lo largo,
       sus secciones la cruzan) y con anchos parecidos.
    3) Cada cadena de 2 o mas secciones se cierra uniendo un extremo con el
       otro a lo largo de toda la cadena.
    Devuelve (areas, bordes, indices_usados).
    """
    from . import unir

    conf = res.conf
    if len(indices) < 4:
        return [], [], set()
    xy = np.array([(res.puntos[i].e, res.puntos[i].n) for i in indices])
    arbol = cKDTree(xy)
    k = min(4, len(xy))
    _, vec = arbol.query(xy, k=k)
    padre = list(range(len(xy)))

    def raiz(i):
        while padre[i] != i:
            padre[i] = padre[padre[i]]
            i = padre[i]
        return i

    for i in range(len(xy)):
        for j in vec[i][1:3]:
            if i in vec[j][1:3] and np.hypot(*(xy[i] - xy[j])) <= conf.ancho_max:
                padre[raiz(i)] = raiz(j)
    grupos = {}
    for i in range(len(xy)):
        grupos.setdefault(raiz(i), []).append(i)
    secciones = []  # (p, q, centro, eje, ancho, miembros)
    for miembros in grupos.values():
        if not 2 <= len(miembros) <= 5:
            continue
        m = xy[miembros]
        dist = np.hypot(*(m[:, None, :] - m[None, :, :]).transpose(2, 0, 1))
        a, b = np.unravel_index(np.argmax(dist), dist.shape)
        ancho = dist[a, b]
        if not conf.ancho_min <= ancho <= conf.ancho_max:
            continue
        p, q = m[a], m[b]
        secciones.append((p, q, (p + q) / 2, (q - p) / ancho, ancho, miembros))
    if len(secciones) < 2:
        return [], [], set()

    centros = np.array([s_[2] for s_ in secciones])
    zona_sec = None
    if getattr(res, "zonas", None) is not None:
        zona_sec = [int(np.bincount(np.asarray(res.zonas)[[indices[m] for m in s_[5]]] + 1).argmax()) - 1
                    for s_ in secciones]
    arbol_c = cKDTree(centros)
    cand = []
    for i, j in arbol_c.query_pairs(conf.separacion_max * 1.25):
        v = centros[j] - centros[i]
        largo = np.hypot(*v)
        u = v / largo
        ci, cj = abs(np.dot(secciones[i][3], u)), abs(np.dot(secciones[j][3], u))
        ratio = max(secciones[i][4], secciones[j][4]) / min(secciones[i][4], secciones[j][4])
        if ci < 0.5 and cj < 0.5 and ratio < 2.5:
            if zona_sec is not None and zona_sec[i] != zona_sec[j]:
                continue  # secciones de manzanas distintas
            if barreras is not None and barreras.cruza(centros[i], centros[j]):
                continue  # cruzaria un limite de propiedad u otra union
            cand.append((largo * (1 + ci + cj + 0.3 * (ratio - 1)), i, j))
    cand.sort()
    grafo = unir._Grafo(centros, barreras or unir.Barreras())
    for costo, i, j in cand:
        if len(grafo.vec[i]) < 2 and len(grafo.vec[j]) < 2 and grafo.raiz(i) != grafo.raiz(j):
            # recto: en cada seccion el giro entre tramos debe ser suave (> 120 grados)
            ok = all(unir._angulo(centros[a], centros[c], centros[b]) >= 120
                     for a, b in ((i, j), (j, i)) for c in grafo.vec[a])
            if ok:
                grafo.vec[i].append(j)
                grafo.vec[j].append(i)
                grafo.padre[grafo.raiz(i)] = grafo.raiz(j)
    areas, bordes, usados = [], [], set()
    for cadena in grafo.cadenas():
        if len(cadena) < 2 or cadena[0] == cadena[-1]:
            continue
        lado1, lado2 = [secciones[cadena[0]][0]], [secciones[cadena[0]][1]]
        for idx in cadena[1:]:
            p, q = secciones[idx][0], secciones[idx][1]
            # orientar la seccion para que los lados no se crucen
            if np.hypot(*(lado1[-1] - p)) + np.hypot(*(lado2[-1] - q)) > \
                    np.hypot(*(lado1[-1] - q)) + np.hypot(*(lado2[-1] - p)):
                p, q = q, p
            lado1.append(p)
            lado2.append(q)
        coords = [tuple(c) for c in lado1] + [tuple(c) for c in lado2[::-1]]
        pol = _poligono_valido(coords)
        if pol.is_empty or pol.area < AREA_MIN:
            continue
        largo = (LineString(lado1).length + LineString(lado2).length) / 2
        anchos = [secciones[i][4] for i in cadena]
        revisar = not Polygon(coords).is_valid
        areas.append(Area(res.codigo, conf, pol, "estaciones", float(np.median(anchos)), largo, revisar))
        bordes += [LineString(lado1), LineString(lado2)]
        for i in cadena:
            usados.update(indices[m] for m in secciones[i][5])
    return areas, bordes, usados


def _unir_sobrantes(res, usados):
    """Los puntos que no entraron en ninguna seccion se unen borde a borde (metodo PETRO)."""
    from . import unir

    resto = [i for i in range(len(res.puntos)) if i not in usados]
    if len(resto) < 2:
        return
    xy = np.array([(res.puntos[i].e, res.puntos[i].n) for i in resto])
    zonas = None if getattr(res, "zonas", None) is None else np.asarray(res.zonas)[resto]
    uniones = unir._candidatos(xy, res.conf, zonas=zonas)
    if res.orto is not None:
        for u in uniones:
            u.apoyo = unir.apoyo_imagen(res.orto, xy[u.i], xy[u.j])
    sub = unir.Resultado(res.codigo, [res.puntos[i] for i in resto], res.conf)
    unir._resolver(sub, xy, uniones, res.orto is not None, res.barreras or unir.Barreras(), res.completar)
    res.uniones = [unir.Union(resto[u.i], resto[u.j], u.largo, u.apoyo, u.costo, u.revisar) for u in sub.uniones]
    res.cadenas = [[resto[i] for i in c] for c in sub.cadenas]


PUNTOS_BORDE_MIN = 3  # puntos topograficos sobre el borde para aceptar un area vista en la foto
TOL_BORDE = 0.4  # m


def pasar_por_puntos(pol, puntos_xy, tol=TOL_BORDE):
    """Ajusta el contorno para que pase exactamente por los puntos topograficos de su borde.

    Cada punto a menos de `tol` del borde se inserta en el contorno (en su posicion a lo largo
    del borde) y se quitan los vertices de la foto que esten a menos de `tol` de ese punto.
    Devuelve (poligono_ajustado, cantidad_de_puntos_en_el_borde).
    """
    anillo = pol.exterior
    largo = anillo.length
    en_borde = [(anillo.project(Point(p)), tuple(p)) for p in puntos_xy if anillo.distance(Point(p)) < tol]
    if not en_borde:
        return pol, 0
    verts = [(anillo.project(Point(c)), c) for c in list(anillo.coords)[:-1]]
    quitar = {i for i, (sv, c) in enumerate(verts)
              for sp, pc in en_borde if np.hypot(c[0] - pc[0], c[1] - pc[1]) < tol}
    nuevos = sorted([v for i, v in enumerate(verts) if i not in quitar] + en_borde, key=lambda t: t[0] % largo)
    ajustado = Polygon([c for _, c in nuevos])
    if not ajustado.is_valid or abs(ajustado.area - pol.area) > 0.3 * pol.area:
        ajustado = pol  # el ajuste deformaria el area: se deja la forma de la foto
    return ajustado, len(en_borde)


def agregar_concreto_visible(met, poligonos, conf, puntos_xy=()):
    """Concreto visto en la foto, coherente con la topografia.

    Solo se acepta (con metrado) si al menos PUNTOS_BORDE_MIN puntos topograficos caen sobre su
    borde; el borde se ajusta para pasar por esos puntos. Lo demas va a met.sin_puntos (revisar).
    """
    ocupado = unary_union([a.poligono for a in met.areas]) if met.areas else Polygon()
    pts = np.asarray(puntos_xy, float).reshape(-1, 2)
    for pol in poligonos:
        resto = pol.difference(ocupado) if not ocupado.is_empty else pol
        for g in getattr(resto, "geoms", [resto]):
            if not isinstance(g, Polygon) or g.area < 5.0:
                continue
            cerca = pts[[g.buffer(TOL_BORDE).contains(Point(p)) for p in pts]] if len(pts) else pts
            ajustado, n = pasar_por_puntos(g, cerca)
            if n >= PUNTOS_BORDE_MIN:
                met.areas.append(Area("CV", conf, ajustado, "foto", revisar=False,
                                      nota=f"concreto visto en la foto; borde ajustado a {n} puntos topograficos"))
            else:
                met.sin_puntos.append(g)
    _numerar(met)


# ---------------------------------------------------------------- edicion a mano: unir y no superponer

TOL_VECINA = 0.30  # m: un area dibujada a menos de esto de otra se considera pegada a ella


def vecinas(pol, lista, tol=TOL_VECINA, excluir=None):
    """Areas de `lista` que tocan, se superponen o quedan a menos de `tol` del poligono."""
    return [a for a in lista if a is not excluir and a.poligono.distance(pol) <= tol]


def _mayor(g):
    if isinstance(g, MultiPolygon):
        return max(g.geoms, key=lambda x: x.area)
    return g if isinstance(g, Polygon) else None


def unir_poligonos(pols, tol=TOL_VECINA):
    """Une poligonos que se tocan o quedan a menos de `tol`, cerrando la ranura entre ellos sin
    engordar el resto del borde. Devuelve un Polygon (el mayor si quedaran separados)."""
    r = tol / 2
    g = unary_union([p.buffer(r, join_style=2, mitre_limit=10) for p in pols]).buffer(-r, join_style=2,
                                                                                       mitre_limit=10)
    g = _mayor(g.simplify(0.005))
    if g is not None and not g.is_valid:
        g = _mayor(g.buffer(0))
    return g


def separar_poligono(pol, otros, tol=TOL_VECINA):
    """El poligono sin lo que pisa de los otros; los vertices a menos de `tol` del borde vecino se
    pegan a el, asi comparten la linea en vez de superponerse o dejar una ranura."""
    from shapely import snap

    if not otros:
        return [pol]
    union = unary_union(otros)
    borde = union.boundary
    # pegar cada vertice cercano al borde vecino (al vertice si esta cerca, si no al lado)
    coords = []
    for x, y in list(pol.exterior.coords)[:-1]:
        q = Point(x, y)
        if borde.distance(q) <= tol:
            q = borde.interpolate(borde.project(q))
        coords.append((q.x, q.y))
    nuevo = Polygon(coords)
    if not nuevo.is_valid:
        nuevo = _mayor(nuevo.buffer(0)) or pol
    nuevo = snap(nuevo, union, 0.01)
    resto = nuevo.difference(union)
    piezas = list(resto.geoms) if isinstance(resto, MultiPolygon) else [resto] if isinstance(resto, Polygon) else []
    return [g for g in piezas if g.area >= AREA_MIN]


def _numerar(met, clave=None):
    """VD - 01, VD - 02... ordenadas de norte a sur y de oeste a este.

    `clave(geom, elemento)`: otro orden (p.ej. por lamina, ver laminas.clave_orden)."""
    if clave is None:
        clave = lambda g, _e: (-round(g.centroid.y / 20), g.centroid.x)  # noqa: E731
    por_prefijo = {}
    for ar in met.areas:
        por_prefijo.setdefault(ar.conf.prefijo or ar.codigo, []).append(ar)
    for pref, lista in por_prefijo.items():
        lista.sort(key=lambda a: clave(a.poligono, a))
        for k, ar in enumerate(lista, 1):
            ar.etiqueta = f"{pref} - {k:02d}"
    por_prefijo = {}
    for ln in met.lineas:
        if ln.conf.prefijo and ln.conf.tipo == "linea" and not ln.borde and not ln.sin_pareja:
            por_prefijo.setdefault(ln.conf.prefijo, []).append(ln)
    for pref, lista in por_prefijo.items():
        lista.sort(key=lambda l: clave(l.linea, l))
        for k, ln in enumerate(lista, 1):
            ln.etiqueta = f"{pref} - {k:02d}"


# Corte lineal (CL): en PETRO es el borde de la vereda a demoler que coincide con el limite de
# propiedad (62 de sus 66 lineas de corte estan sobre LINEA DE LOTE y sobre la vereda).
CORTE_TOL = 0.10  # m: borde del area a esta distancia del limite de propiedad = corte
CORTE_MIN = 1.0  # m: tramos de corte mas cortos que esto no se dibujan
CORTE_HUECO = 1.5  # m: tramos de corte separados por menos que esto forman una sola linea


def conf_corte():
    from .unir import Codigo

    return Codigo(capa="LINEA CORTE", tipo="linea", prefijo="CL", nombre="Corte lineal", color=242,
                  tipo_linea="ACAD_ISO10W100", separacion_max=0)


def cortes_lineales(met, base=None, resultados=()):
    """Agrega a met.lineas las lineas de corte (CL) de las areas cuyo codigo tiene corte_lineal."""
    met.lineas = [l for l in met.lineas if l.codigo != "CL"]
    limites = list(base.limites) if base is not None and not base.vacio else []
    refs = set()
    for a in met.areas:
        if getattr(a.conf, "corte_lineal", False):
            refs.update(a.conf.referencia or [])
    for r in resultados:
        if r.codigo in refs and r.conf.tipo == "linea":
            for cad in r.cadenas:
                if len(cad) >= 2:
                    limites.append(LineString(r.xy(cad)))
    if not limites:
        return []
    zona = unary_union([l.buffer(CORTE_TOL) for l in limites])
    arbol = STRtree(limites)

    def paralelo(a, b):
        """El tramo a-b corre a lo largo del limite (no es el extremo de la vereda que lo cruza)."""
        m = Point((a[0] + b[0]) / 2, (a[1] + b[1]) / 2)
        lim = limites[int(arbol.nearest(m))]
        d = lim.project(m)
        p0, p1 = lim.interpolate(max(d - 0.05, 0)), lim.interpolate(min(d + 0.05, lim.length))
        u = np.array([p1.x - p0.x, p1.y - p0.y])
        v = np.array([b[0] - a[0], b[1] - a[1]])
        nu, nv = np.hypot(*u), np.hypot(*v)
        return nu > 1e-9 and nv > 1e-9 and abs(u @ v) / (nu * nv) > 0.87  # menos de 30 grados

    # Cada tramo de borde paralelo al limite se proyecta sobre ese limite; los tramos seguidos (con
    # huecos de menos de CORTE_HUECO m, p.ej. entre dos veredas vecinas) forman una sola linea de corte
    # que va por el limite de propiedad, como en PETRO.
    intervalos = {}
    for a in met.areas:
        if not getattr(a.conf, "corte_lineal", False):
            continue
        c = list(a.poligono.exterior.coords)
        for p, q in zip(c[:-1], c[1:]):
            g = LineString([p, q]).intersection(zona)
            for parte in getattr(g, "geoms", [g]):
                if isinstance(parte, LineString) and parte.length > 0.01:
                    pc = list(parte.coords)
                    if paralelo(pc[0], pc[-1]):
                        m = parte.interpolate(0.5, normalized=True)
                        k = int(arbol.nearest(m))
                        t0, t1 = sorted((limites[k].project(Point(pc[0])), limites[k].project(Point(pc[-1]))))
                        intervalos.setdefault(k, []).append([t0, t1])
    conf = conf_corte()
    nuevas = []
    for k, ivs in intervalos.items():
        ivs.sort()
        unidos = [ivs[0]]
        for t0, t1 in ivs[1:]:
            if t0 - unidos[-1][1] <= CORTE_HUECO:
                unidos[-1][1] = max(unidos[-1][1], t1)
            else:
                unidos.append([t0, t1])
        for t0, t1 in unidos:
            if t1 - t0 >= CORTE_MIN:
                ln = substring(limites[k], t0, t1)
                if isinstance(ln, LineString) and ln.length >= CORTE_MIN:
                    nuevas.append(Linea("CL", conf, ln, dibujar=True))
    met.lineas += nuevas
    return nuevas


def punto_etiqueta(pol):
    """Punto dentro del poligono para poner el texto."""
    p = pol.representative_point()
    return p.x, p.y


def guardar_metrado(met, ruta_xlsx, ruta_csv=None, laminas=None):
    """Planilla de metrados (Excel) y CSV. Las columnas nuevas (perimetro, lamina) van al final."""
    filas = []
    for ar in met.areas:
        x, y = punto_etiqueta(ar.poligono)
        filas.append([ar.etiqueta, ar.conf.nombre or ar.codigo, ar.conf.capa_area, "m2", round(ar.area, 2),
                      round(ar.largo, 2), round(ar.ancho, 2), "SI" if ar.revisar else "", round(x, 3), round(y, 3),
                      ar.nota, round(ar.poligono.exterior.length, 2), ar.lamina, "; ".join(ar.reglas)])
    for ln in met.lineas:
        if not ln.etiqueta:
            continue
        x, y = ln.linea.interpolate(0.5, normalized=True).coords[0]
        filas.append([ln.etiqueta, ln.conf.nombre or ln.codigo, ln.conf.capa, "m", round(ln.linea.length, 2),
                      round(ln.linea.length, 2), "", "", round(x, 3), round(y, 3), "", "", ln.lamina, ""])
    cab = ["Etiqueta", "Elemento", "Capa", "Unidad", "Metrado", "Largo (m)", "Ancho medio (m)", "Revisar",
           "Este", "Norte", "Observacion", "Perimetro (m)", "Lamina", "Reglas aplicadas"]
    resumen = {}
    for f in filas:
        clave = (f[0].split(" - ")[0], f[1], f[3])
        n, tot = resumen.get(clave, (0, 0.0))
        resumen[clave] = (n + 1, tot + f[4])
    if ruta_csv:
        import csv

        with open(ruta_csv, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh, delimiter=";")
            w.writerow(cab)
            w.writerows(filas)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return resumen
    wb = Workbook()
    hoja = wb.active
    hoja.title = "Resumen"
    hoja.append(["Codigo", "Elemento", "Unidad", "Cantidad", "Metrado total"])
    for (pref, nombre, und), (n, tot) in sorted(resumen.items()):
        hoja.append([pref, nombre, und, n, round(tot, 2)])
    det = wb.create_sheet("Detalle")
    det.append(cab)
    for f in filas:
        det.append(f)
    hojas = [hoja, det]
    if any(f[12] for f in filas):
        por = wb.create_sheet("Por lamina")
        por.append(["Lamina", "Codigo", "Elemento", "Unidad", "Cantidad", "Metrado total", "Perimetro total (m)"])
        tot_l = {}
        for f in filas:
            k = (f[12], f[0].split(" - ")[0], f[1], f[3])
            n, m, p = tot_l.get(k, (0, 0.0, 0.0))
            tot_l[k] = (n + 1, m + f[4], p + (f[11] or 0.0))
        for (lam, pref, nombre, und), (n, m, p) in sorted(tot_l.items()):
            por.append([lam, pref, nombre, und, n, round(m, 2), round(p, 2) if und == "m2" else ""])
        hojas.append(por)
    amarillo = PatternFill("solid", fgColor="FFF2CC")
    for h in hojas:
        for c in h[1]:
            c.font = Font(bold=True)
        for col in h.columns:
            h.column_dimensions[col[0].column_letter].width = max(10, max(len(str(c.value or "")) for c in col) + 2)
    for fila in det.iter_rows(min_row=2):
        if fila[7].value == "SI":
            for c in fila:
                c.fill = amarillo
    wb.save(ruta_xlsx)
    return resumen
